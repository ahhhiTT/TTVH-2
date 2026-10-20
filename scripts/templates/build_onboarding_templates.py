"""Build the onboarding templates served from public/templates/ (structure only, no brand data).

    python3 -m pip install openpyxl
    python3 scripts/templates/build_onboarding_templates.py

Sources of the structure (Director, 2026-10-20): master-data sheets agreed for the weekly
report, and the JBP layout of a 6-month TikTok Shop plan (sections A/B/C).
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = os.path.join(os.path.dirname(__file__), "..", "..", "public", "templates")
os.makedirs(OUT, exist_ok=True)

HEAD = PatternFill("solid", fgColor="171717")
SEC = PatternFill("solid", fgColor="E8E8EC")
FILL = PatternFill("solid", fgColor="FFF4D6")
EX = Font(italic=True, color="7A7A7A")
thin = Side(style="thin", color="D4D4D8")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")


def table(wb, title, headers, widths, examples=(), dropdowns=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(1, i)
        c.fill, c.font, c.alignment, c.border = HEAD, Font(bold=True, color="FFFFFF"), WRAP, BOX
    ws.freeze_panes = "A2"
    for ex in examples:
        ws.append(list(ex))
        for c in ws[ws.max_row]:
            c.font, c.alignment, c.border = EX, WRAP, BOX
    for col, opts in (dropdowns or {}).items():
        dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}2000")
    return ws


def guide(ws, rows):
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 100
    for a, b in rows:
        ws.append([a, b])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        ws.cell(ws.max_row, 2).alignment = WRAP
    ws["A1"].font = Font(bold=True, size=14)


# ------------------------------------------------------------------ master data
wb = Workbook()
guide(wb.active, [
    ("MẪU DANH MỤC ONBOARDING", ""),
    ("Dùng khi", "Brand hoặc gian hàng mới onboard. Nhân sự phụ trách điền đủ các sheet và nộp trên Brand Hub."),
    ("Dòng chữ xám nghiêng", "Ví dụ, số và tên đều giả. Xóa khi điền thật."),
    ("Định dạng", "Ngày: YYYY-MM-DD. Tiền: số nguyên VND. Không để trống cột mã."),
    ("DM_Gian_hang", "Mỗi gian hàng một dòng: gian hàng trong hệ thống ứng với shop nào trên sàn."),
    ("DM_Tai_khoan_kenh", "Tài khoản TikTok/kênh của brand (self channel): kênh chính, kênh phụ. Dùng để tách GMV self channel và affiliate."),
    ("DM_San_pham", "File sản phẩm chuẩn hóa: mỗi SKU người bán một dòng, gắn dòng sản phẩm và danh mục. Combo ghi thành phần."),
    ("DM_Creator_Tier", "Creator đã book và tier. Có thể lấy từ Base Booking."),
])
wb.active.title = "Huong_dan"
table(wb, "DM_Gian_hang",
      ["Mã gian hàng (hệ thống)", "Brand", "Sàn", "Tên shop trên Seller Center", "Shop ID / Shop code", "Nhân sự phụ trách (email)", "Hiệu lực từ", "Hiệu lực đến", "Ghi chú"],
      [18, 18, 14, 30, 22, 28, 14, 14, 30],
      [("s00", "Brand Mẫu", "TikTok Shop", "Brand Mẫu Official", "VD000000", "nhansu@example.com", "2026-10-01", "", "Ví dụ")],
      {"C": ["TikTok Shop", "Shopee", "Lazada", "Website"]})
table(wb, "DM_Tai_khoan_kenh",
      ["Mã gian hàng", "Tên tài khoản", "@handle", "User ID (nếu có)", "Loại kênh", "Hiệu lực từ", "Hiệu lực đến", "Ghi chú"],
      [16, 26, 22, 22, 16, 14, 14, 30],
      [("s00", "Brand Mẫu Official", "@brandmau.official", "", "Kênh chính", "2026-10-01", "", "Ví dụ")],
      {"E": ["Kênh chính", "Kênh phụ 1", "Kênh phụ 2", "Kênh phụ 3"]})
table(wb, "DM_San_pham",
      ["Brand", "Mã SKU người bán", "Product ID trên sàn", "SKU ID trên sàn", "Tên sản phẩm chuẩn", "Tên trên sàn", "Dòng sản phẩm",
       "Danh mục sản phẩm", "Loại", "Thành phần combo (Mã SKU x SL; ...)", "Giá niêm yết", "Trạng thái", "Hiệu lực từ", "Ghi chú"],
      [16, 16, 20, 20, 32, 32, 18, 18, 12, 32, 14, 12, 14, 24],
      [("Brand Mẫu", "SP00001", "1700000000000000001", "1700000000000000101", "Sản phẩm X 300ml", "Sản phẩm X chính hãng 300ml", "Dòng X", "Chăm sóc cơ thể", "Đơn lẻ", "", "199000", "Đang bán", "2026-10-01", "Ví dụ"),
       ("Brand Mẫu", "CB00001", "1700000000000000002", "1700000000000000102", "Combo 2 Sản phẩm X", "Combo 2 chai Sản phẩm X", "Dòng X", "Chăm sóc cơ thể", "Combo", "SP00001 x 2", "379000", "Đang bán", "2026-10-01", "Ví dụ")],
      {"I": ["Đơn lẻ", "Combo", "Quà tặng"], "L": ["Đang bán", "Ngừng bán"]})
table(wb, "DM_Creator_Tier",
      ["Brand", "Tên creator", "@handle", "Creator ID (nếu có)", "Tier", "Tệp kênh", "MCN / Agency", "Hiệu lực từ", "Ghi chú"],
      [16, 24, 22, 20, 16, 18, 20, 14, 26],
      [("Brand Mẫu", "Creator Mẫu", "@creator.mau", "", "KL3", "Mẹ bé", "", "2026-10-01", "Ví dụ")],
      {"E": ["Freecast", "TAP UpAffiliate", "KL1", "KL2", "KL3", "KL4", "KL5", "KL6", "KL7"]})
wb.save(os.path.join(OUT, "mau-danh-muc-onboarding.xlsx"))

# ------------------------------------------------------------------ JBP
wb = Workbook()
ws = wb.active
ws.title = "JBP"
MONTHS = ["Tháng 1", "Tháng 2", "Tháng 3", "Tháng 4", "Tháng 5", "Tháng 6", "Tổng giai đoạn", "Ước thực hiện kỳ trước", "Tổng năm cập nhật", "JBP đầu năm", "GAP"]
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 44
for i in range(3, 3 + len(MONTHS)):
    ws.column_dimensions[get_column_letter(i)].width = 15


def sec(title):
    ws.append([title])
    for c in ws[ws.max_row][:2 + len(MONTHS)]:
        c.fill = SEC
    ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)


def header(cols):
    ws.append(cols)
    for c in ws[ws.max_row]:
        if c.value is not None:
            c.fill, c.font, c.alignment = HEAD, Font(bold=True, color="FFFFFF"), WRAP


def rows(group, items, fill=True):
    for i, it in enumerate(items):
        ws.append([group if i == 0 else None, it] + ([None] * len(MONTHS)))
        if fill:
            for c in ws[ws.max_row][2:2 + len(MONTHS)]:
                c.fill = FILL


ws.append(["KẾ HOẠCH KINH DOANH CHUNG (JBP)", "Brand: ...    Gian hàng: ...    Giai đoạn: ..."])
ws["A1"].font = Font(bold=True, size=14)
ws.append(["Ô màu vàng là chỗ điền. Đổi tên cột Tháng 1..6 thành tháng thật (vd T7.2026). Giai đoạn có thể là năm, quý hoặc 6 tháng."])
sec("A. PLAN SỐ TỔNG QUAN")
header(["Hạng mục", "Chỉ số"] + MONTHS)
rows("Chỉ số chung", ["Tăng trưởng MoM", "GMV nhà bán", "NMV nhà bán", "Số đơn hàng hiệu quả", "Số đơn hàng", "Lượt truy cập gian hàng",
                      "Tỉ lệ chuyển đổi (CR)", "Giá trung bình/đơn hàng (AOV)", "Tỷ lệ hủy", "Tổng chi phí thúc đẩy", "% Chi phí thúc đẩy/NMV (CIR)"])
rows("Doanh thu", ["GAP (GMV sàn/GMV nhà bán)", "GMV sàn", "Doanh thu Video", "Video self channel", "Video affiliate",
                   "Doanh thu Livestream", "Livestream self channel", "Livestream affiliate", "Doanh thu trang sản phẩm"])
rows("Chi phí thúc đẩy", ["Tổng chi phí onsite", "Chi phí booking video affiliate", "Chi phí video self channel", "Chi phí Ads (gồm thuế sàn)",
                          "Chi phí booking livestream affiliate", "Chi phí livestream self channel", "Chi phí tiếp thị liên kết",
                          "Chi phí gói Marketing Solution", "Chi phí gói dịch vụ (voucher, hoàn xu, freeship extra...)", "Chi phí khác",
                          "Tổng chi phí offsite (gross)", "Tổng chi phí extra (branding nội sàn/ngoại sàn...)"])
rows("Chi phí sản phẩm", ["Tổng chi phí sản phẩm", "Chi phí chiết khấu", "Chi phí quà tặng", "Chi phí sample", "Số lượng sample dự kiến"])
rows("Chi phí vận hành", ["Tổng chi phí sàn", "Chi phí giao dịch", "Chi phí hoa hồng cố định", "Chi phí hạ tầng (xử lý đơn hàng)", "Chi phí vận chuyển nhà bán chịu"])
rows("Chi phí UpBase", ["Tổng chi phí UpBase", "Chi phí fix", "Chi phí hoa hồng", "Chi phí fulfillment"])
rows("Chi phí COGS", ["Chi phí COGS"])
sec("B. BREAKDOWN")
header(["Nhóm", "Dòng sản phẩm / hạng mục"] + MONTHS)
rows("1. GMV mục tiêu theo dòng sản phẩm", ["[Dòng sản phẩm 1]", "[Dòng sản phẩm 2]", "[Dòng sản phẩm 3]", "Khác"])
rows("2. Chi phí mục tiêu (net)", ["NMV gian hàng", "Tổng chi phí marketing", "% Chi phí marketing", "Chi phí booking video affiliate",
                                     "Chi phí video self channel", "Chi phí Ads", "Chi phí booking livestream affiliate",
                                     "Chi phí livestream self channel", "Chi phí tiếp thị liên kết", "Chi phí khác", "Ngân sách branding đề xuất"])
rows("3. Booking", ["Chi phí booking video affiliate (Performance) theo dòng sản phẩm", "Số video booking (KL1-7) theo dòng sản phẩm",
                    "Số video TAP / free of cast / Aff Hub theo dòng sản phẩm", "Ngân sách top up booking KOL key branding"])
rows("4. Ads", ["% GMV Ads / GMV nhà bán", "GMV Ads (toàn gian hàng)", "ROAS gian hàng", "Chi phí Ads thực chạy", "GMV PGM", "ROAS PGM", "Cost PGM",
                "GMV LGM", "ROAS LGM", "Cost LGM", "Consideration: size, cost, CPA", "Branding: view, follow, cost",
                "PGM theo dòng sản phẩm: GMV, ROAS, cost", "LGM kênh / LGM affiliate: GMV, ROAS, cost", "Số phiên live affiliate có Ads / không Ads"])
sec("C. CHIẾN LƯỢC TRIỂN KHAI")
header(["Hạng mục", "Nội dung cần có"] + MONTHS[:6])
STRAT = [
    ("1. Chiến lược sản phẩm", ["Campaign timeline: campaign sàn, campaign brand/sản phẩm/gian hàng",
                                "Product roadmap: sản phẩm nào key focus / duy trì / bán tự nhiên theo từng tháng",
                                "Vai trò nhóm sản phẩm: Entry, Expansion, Key"]),
    ("2. Chiến lược commercial", ["Nhóm sản phẩm (ra mắt, key focus, duy trì, combo): giá niêm yết, giá daily, quà daily"]),
    ("3. Chiến lược content", ["Campaign brand", "Always-on content", "Content direction", "Affiliate", "Content kênh"]),
    ("4. Chiến lược kênh và Ads", ["PGM sản phẩm cũ / mới", "LGM kênh", "LGM affiliate", "Branding Ads", "Chiến lược chung"]),
    ("5. Key winning / activities", ["Hoạt động chính từng tháng"]),
    ("6. Quản trị rủi ro", ["Rủi ro và giải pháp: campaign marketing của brand, booking KOL/KOC, phản hồi thị trường, CR/traffic"]),
]
for g, items in STRAT:
    rows(g, items)
wb.save(os.path.join(OUT, "mau-jbp.xlsx"))
print("written to", os.path.abspath(OUT))
